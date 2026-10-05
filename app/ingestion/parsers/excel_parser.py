"""Parser de la planilla de precios de carnes y pescados."""

import logging
from pathlib import Path
import openpyxl
from openpyxl.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from app.ingestion.normalization import normalize_name
from app.ingestion.parsers.common import PriceParser, parse_price
from app.ingestion.parsers.schemas import ParsedPrice, PriceParseResult

logger = logging.getLogger(__name__)

SECTION_TITLES = ("Carnicería", "Pescadería")


class ExcelPriceParser(PriceParser):
    def __init__(self, path: Path):
        super().__init__(path)
        workbook = openpyxl.load_workbook(path, data_only=True)
        self._sheet: Worksheet = workbook.active
        # Se agrega la hoja al origen para que el rechazo diga dónde mirar.
        self._source = f"{path.name}:{self._sheet.title}"

    def parse(self) -> PriceParseResult:
        result = PriceParseResult()
        for title in SECTION_TITLES:
            anchor = self._find_cell(title)
            if anchor is None:
                self._reject(result, None, title, "no se encontró la sección en la planilla")
                continue
            self._parse_section(anchor, result)
        return result

    def _find_cell(self, text: str) -> Cell | None:
        target = normalize_name(text)
        for row in self._sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and normalize_name(cell.value) == target:
                    return cell
        return None

    def _parse_section(self, title_cell: Cell, result: PriceParseResult) -> None:
        name_column = title_cell.column
        price_column = name_column + 1
        found_rows = False

        for row in range(title_cell.row + 1, self._sheet.max_row + 1):
            name_cell = self._sheet.cell(row=row, column=name_column)
            raw_name = name_cell.value
            raw_price = self._sheet.cell(row=row, column=price_column).value

            if raw_name is None and raw_price is None:
                if found_rows:
                    break
                continue

            if self._is_header(raw_price):
                continue

            # Los subtítulos ("Carne de Cerdo", "Pollo") ocupan nombre y precio en una celda combinada. Así se distinguen de un producto sin precio.
            if name_cell.coordinate in self._sheet.merged_cells:
                logger.debug("Subtítulo ignorado en fila %d: %r", row, raw_name)
                continue

            found_rows = True
            location = f"fila {row}"
            content = f"nombre={raw_name!r}, precio={raw_price!r}"
            name = str(raw_name).strip() if raw_name is not None else ""

            if not name:
                self._reject(result, location, content, "fila sin nombre")
                continue
            if raw_price is None:
                self._reject(result, location, content, "fila sin precio")
                continue
            try:
                price = parse_price(raw_price)
            except ValueError as exc:
                self._reject(result, location, content, str(exc))
                continue

            result.prices.append(ParsedPrice(self._source, location, name, price))

    @staticmethod
    def _is_header(raw_price: object) -> bool:
        return isinstance(raw_price, str) and "precio" in normalize_name(raw_price)
